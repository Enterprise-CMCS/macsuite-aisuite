#!/usr/bin/env python3

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from common.utils.logger import log
from search.excel_process.process_excel_with_rag import (
    ANALYSIS_COLUMNS,
    ANALYSIS_SHEET,
    STATUS_COL,
    STATUS_WORDING,
    CRTWorkbook,
    _text,
    default_folder,
    find_workbooks,
    output_path,
)

DESCRIPTION = "Compare AI-reviewed CRT workbooks against their originals, tab by tab."

BUCKETS = ("Met", "Not Met", "Unsure", "Blank", "Unrecognized")

REPORT_NAME = "comparison_report.xlsx"

SUMMARY_HEADER = ["Workbook", "Tab", "Status", "Original Count", "Reviewed Count", "Delta"]
DIFF_HEADER = ["Workbook", "Tab", "Row", "Item", "Original Status", "Reviewed Status"]

PPT_HEADER = ["Submission Tab", "Total Requirements", "Unanswered (Original)", "Met",
              "Not Met / Unclear", "Recommendation Accuracy", "Citation Accuracy"]

# Column position of "Quotes Verified" within ANALYSIS_COLUMNS, resolved by name so this
# doesn't silently drift if process_excel_with_rag.py reorders the RAG Analysis sheet.
_ANALYSIS_SHEET_COL = 1
_ANALYSIS_ROW_COL = 2
_ANALYSIS_QUOTES_VERIFIED_COL = [header for header, _ in ANALYSIS_COLUMNS].index("Quotes Verified") + 1

THIN_BORDER = Border(*(Side(style="thin"),) * 4)


def normalize_status(value):
    text = _text(value)
    if not text:
        return "Blank"

    lowered = text.lower()
    for canonical, variants in STATUS_WORDING.items():
        if lowered in variants:
            return canonical.title()
    return "Unrecognized"


def empty_counts():
    return {bucket: 0 for bucket in BUCKETS}


def collect_pairs(folder):
    originals = [path for path in find_workbooks(folder) if path.name != REPORT_NAME]
    pairs, skipped = [], []
    for original in originals:
        reviewed = output_path(original)
        if reviewed.exists():
            pairs.append((original, reviewed))
        else:
            skipped.append(original)
    return pairs, skipped


def load_quotes_verified(reviewed):
    """Map (tab title, row) -> "Yes"/"No"/"" from the reviewed workbook's RAG Analysis sheet.

    "" means the row has no citation to verify (no evidence was quoted), so it's excluded
    from the citation-accuracy denominator rather than counted as a miss.
    """
    verified = {}
    if ANALYSIS_SHEET not in reviewed.workbook.sheetnames:
        return verified

    sheet = reviewed.workbook[ANALYSIS_SHEET]
    for row in sheet.iter_rows(min_row=2):
        title = _text(row[_ANALYSIS_SHEET_COL - 1].value)
        row_number = row[_ANALYSIS_ROW_COL - 1].value
        if not title or row_number is None:
            continue
        verified[(title, int(row_number))] = _text(row[_ANALYSIS_QUOTES_VERIFIED_COL - 1].value)

    return verified


def compare_tab(title, requirement_rows, original_sheet, reviewed_sheet, quotes_verified):
    orig_counts, rev_counts, diffs = empty_counts(), empty_counts(), []
    met_total = met_agree = citation_total = citation_yes = 0

    for entry in requirement_rows:
        row = entry["row"]
        orig_bucket = normalize_status(original_sheet.cell(row=row, column=STATUS_COL).value)
        rev_bucket = normalize_status(reviewed_sheet.cell(row=row, column=STATUS_COL).value)
        orig_counts[orig_bucket] += 1
        rev_counts[rev_bucket] += 1
        if orig_bucket != rev_bucket:
            diffs.append({"row": row, "item": entry["item"], "original": orig_bucket, "reviewed": rev_bucket})

        # "Total Requirements" is scoped to what the original analyst marked Met - that's the
        # population the PPT summary grades the AI review against.
        if orig_bucket == "Met":
            met_total += 1
            if rev_bucket == "Met":
                met_agree += 1

            verified = quotes_verified.get((title, row), "")
            if verified in ("Yes", "No"):
                citation_total += 1
                if verified == "Yes":
                    citation_yes += 1

    return {
        "title": title,
        "orig_counts": orig_counts,
        "rev_counts": rev_counts,
        "diffs": diffs,
        "met_total": met_total,
        "met_agree": met_agree,
        "met_mismatch": met_total - met_agree,
        "citation_total": citation_total,
        "citation_yes": citation_yes,
    }


def compare_pair(original_path, reviewed_path):
    original = CRTWorkbook(original_path)
    reviewed = CRTWorkbook(reviewed_path)

    orig_titles = {sheet.title for sheet in original.requirement_sheets}
    rev_titles = {sheet.title for sheet in reviewed.requirement_sheets}
    quotes_verified = load_quotes_verified(reviewed)

    tabs = []
    for title in sorted(orig_titles & rev_titles):
        rows = original.requirements(sheet_names=[title])
        tabs.append(compare_tab(title, rows, original.workbook[title], reviewed.workbook[title],
                                 quotes_verified))

    totals_orig, totals_rev = empty_counts(), empty_counts()
    for tab in tabs:
        for bucket in BUCKETS:
            totals_orig[bucket] += tab["orig_counts"][bucket]
            totals_rev[bucket] += tab["rev_counts"][bucket]

    log.info(f"compare_pair() {original_path.name}: compared {len(tabs)} tab(s) against {reviewed_path.name}")

    return {
        "original_name": original_path.name,
        "reviewed_name": reviewed_path.name,
        "tabs": tabs,
        "only_original": sorted(orig_titles - rev_titles),
        "only_reviewed": sorted(rev_titles - orig_titles),
        "totals": {"orig": totals_orig, "rev": totals_rev},
    }


def _print_counts_table(orig_counts, rev_counts):
    print(f"    {'Status':<14}{'Original':>10}{'Reviewed':>10}{'Delta':>8}")
    for bucket in BUCKETS:
        o, r = orig_counts[bucket], rev_counts[bucket]
        if o == 0 and r == 0:
            continue
        delta = r - o
        sign = f"+{delta}" if delta > 0 else str(delta)
        print(f"    {bucket:<14}{o:>10}{r:>10}{sign:>8}")


def print_comparison(result):
    print(f"\n{result['original_name']}  vs  {result['reviewed_name']}")
    if result["only_original"]:
        print(f"  Sheets only in original (skipped): {', '.join(result['only_original'])}")
    if result["only_reviewed"]:
        print(f"  Sheets only in reviewed (skipped): {', '.join(result['only_reviewed'])}")

    for tab in result["tabs"]:
        print(f"\n  {tab['title']}")
        _print_counts_table(tab["orig_counts"], tab["rev_counts"])

        if tab["diffs"]:
            print(f"    {len(tab['diffs'])} row(s) changed status:")
            for diff in tab["diffs"]:
                print(f"      Row {diff['row']} ({diff['item']}): {diff['original']} -> {diff['reviewed']}")
        else:
            print("    No row-level status changes")

    print(f"\n  TOTAL across {len(result['tabs'])} tab(s)")
    _print_counts_table(result["totals"]["orig"], result["totals"]["rev"])


def _autofit(sheet):
    sheet.freeze_panes = "A2"
    for column_cells in sheet.columns:
        longest = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=0)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(longest + 2, 10), 60)


def _pct(numerator, denominator):
    return f"{numerator / denominator:.0%}" if denominator else "N/A"


def build_ppt_summary(workbook, results):
    multi = len(results) > 1
    header = (["Workbook"] + PPT_HEADER) if multi else PPT_HEADER

    sheet = workbook.active
    sheet.title = "PPT Summary"
    sheet.append(header)

    for result in results:
        prefix = [result["original_name"]] if multi else []
        for tab in result["tabs"]:
            if tab["met_total"] == 0:
                # Nothing in this tab was marked Met by the original analyst, so there's no
                # AI-review population to grade here.
                continue
            sheet.append(prefix + [
                tab["title"],
                tab["met_total"],
                tab["orig_counts"]["Blank"],
                tab["met_agree"],
                tab["met_mismatch"],
                _pct(tab["met_agree"], tab["met_total"]),
                _pct(tab["citation_yes"], tab["citation_total"]),
            ])

    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=len(header)):
        for cell in row:
            cell.border = THIN_BORDER
            if row[0].row == 1:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")
            elif cell.column > (2 if multi else 1):
                cell.alignment = Alignment(horizontal="center")

    _autofit(sheet)
    return sheet


def build_report(results, folder, report_name=REPORT_NAME):
    workbook = Workbook()

    build_ppt_summary(workbook, results)

    summary = workbook.create_sheet("Summary")
    summary.append(SUMMARY_HEADER)

    diffs_sheet = workbook.create_sheet("Row Diffs")
    diffs_sheet.append(DIFF_HEADER)

    for sheet in (summary, diffs_sheet):
        for column in range(1, sheet.max_column + 1):
            sheet.cell(row=1, column=column).font = Font(bold=True)

    for result in results:
        name = result["original_name"]
        for tab in result["tabs"]:
            for bucket in BUCKETS:
                o, r = tab["orig_counts"][bucket], tab["rev_counts"][bucket]
                if o == 0 and r == 0:
                    continue
                summary.append([name, tab["title"], bucket, o, r, r - o])
            for diff in tab["diffs"]:
                diffs_sheet.append([name, tab["title"], diff["row"], diff["item"],
                                    diff["original"], diff["reviewed"]])

        for bucket in BUCKETS:
            o, r = result["totals"]["orig"][bucket], result["totals"]["rev"][bucket]
            if o == 0 and r == 0:
                continue
            summary.append([name, "TOTAL", bucket, o, r, r - o])

    _autofit(summary)
    _autofit(diffs_sheet)

    report_path = folder / report_name
    workbook.save(report_path)
    log.info(f"build_report() Wrote comparison report to {report_path}")
    return report_path


def run(args):
    folder = Path(args.folder) if args.folder else default_folder()
    if not folder.is_dir():
        print(f"Folder not found: {folder}")
        return 1

    pairs, skipped = collect_pairs(folder)
    if skipped:
        print("No .reviewed workbook found for (skipped):")
        for path in skipped:
            print(f"  {path.name}")

    if not pairs:
        print(f"No original/.reviewed workbook pairs found in {folder}")
        return 1

    results = [compare_pair(original, reviewed) for original, reviewed in pairs]
    for result in results:
        print_comparison(result)

    report_path = build_report(results, folder)
    print(f"\nWritten comparison report to {report_path}")
    return 0


def main():
    sys.stdout.reconfigure(errors="replace")

    parser = argparse.ArgumentParser(description=DESCRIPTION)
    parser.add_argument("--folder", help="Folder holding the workbook pairs (default: output_excel/)")
    args = parser.parse_args()

    return run(args)


if __name__ == "__main__":
    sys.exit(main())
