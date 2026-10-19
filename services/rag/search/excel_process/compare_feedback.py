#!/usr/bin/env python3

import argparse
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from common.utils.logger import log
from search.excel_process.compare_review import THIN_BORDER, normalize_status
from search.excel_process.process_excel_with_rag import (
    COMMENTS_COL,
    FIRST_DATA_ROW,
    FOLLOW_UP_COL,
    HEADER_ROW,
    ITEM_COL,
    LEGAL_CITE_COL,
    REQUIREMENT_COL,
    STATUS_COL,
    WHERE_FOUND_COL,
    _text,
    default_folder,
    saveable,
)

DESCRIPTION = ("Compare reviewer feedback on an earlier AI-reviewed CRT workbook against a newer "
               "AI-reviewed output of the same workbook.")

REPORT_NAME = "feedback_comparison_report.xlsx"

HEADER = ["Sheet", "Row", "Item", "Legal Cite", "Requirement",
          "Old: Status", "Old: Where Found", "Old: Follow-up", "Old: AI Comments",
          "Reviewer", "Reviewer Feedback", "Expected Status (from feedback)",
          "New: Status", "New: Where Found", "New: Follow-up", "New: AI Comments",
          "Status Changed?", "Outcome vs Feedback"]

WIDTHS = [24, 6, 12, 18, 50, 12, 24, 40, 50, 10, 70, 18, 12, 24, 40, 50, 10, 22]

# Reviewer columns sit to the right of the CRT's own columns with a header like "John's Comments"
# or "Lynn's comments (AI is using pdf page #s ...)" on the CRT header row.
REVIEWER_HEADER = re.compile(r"^\s*(\w+)'s comments", re.IGNORECASE)

# The reviewed workbook has helper columns well past the CRT block (list values, flags), but
# reviewer columns have so far been within the first ~60.
MAX_SCAN_COL = 80
MAX_SCAN_ROW = 2000

# Ordered: first match wins. Each entry maps a phrase the reviewers actually use to the status
# that phrase implies. "KEEP" means the reviewer endorsed the old AI status.
FEEDBACK_RULES = [
    (re.compile(r"should be unmet|unmet rather than", re.I), "Not Met"),
    (re.compile(r"should be \"?n/?a\"?|be \"n/a\"|likely be \"n/a\"", re.I), "N/A"),
    (re.compile(r"clearly \"?met\"?|consider (this|it)( item)?( to be)? \"?met\"?|would meet|"
                r"shoul?d? be \"?met\"?|clearly meet|explicitly complies|have just missed|"
                r"missed the provision", re.I), "Met"),
    (re.compile(r"agree with ai|^agreed|good ai response|ai response is good|"
                r"determination \(\w+\) is good|correct response", re.I), "KEEP"),
]

# Feedback where the reviewer hedges or the expected status needs reading the whole comment.
# Keyed by (sheet, item) so it survives row shifts; each carries the reason shown in the report.
EXPECTED_OVERRIDES = {
    ("B. Enrollment & Disenrollment", "I.B.1.03"): (
        "Unsure", "Contract covers sexual orientation but not sex; reviewer would clarify with the State."),
    ("B. Enrollment & Disenrollment", "I.B.2.02"): (
        "Unsure", "AI missed excerpts, but reviewer says follow-up with the State is defensible."),
    ("B. Enrollment & Disenrollment", "I.B.3.01"): (
        "Unsure", "Reviewer is unsure; says the AI's response (Unsure + follow-up) is okay."),
    ("B. Enrollment & Disenrollment", "I.B.6.01"): (
        "Met or Unsure", "\"Arguably okay\" (Met), but a discrepancy worth following up with the State."),
    ("B. Enrollment & Disenrollment", "I.B.7.01"): (
        "Met or Unsure", "\"Arguably this is met\", but \"not a bad idea to follow up with the State\"."),
    ("C. Beneficiary Notification", "I.C.1.09"): (
        "Unsure", "\"Good AI response\"; reasonable to clarify with the State, though a close call."),
    ("C. Beneficiary Notification", "I.C.1.12"): (
        "Met", "\"I would consider this to be met for the same reasons\" as the earlier item."),
    ("C. Beneficiary Notification", "I.C.4.01–08"): (
        "Unsure", "Hard copy requirement met, but a clarification question to the State is reasonable."),
    ("D. Payment", "I.D.1.02"): (
        "Unsure", "State as payer is clear, but rate retention needs an inference to call it met."),
    ("D. Payment", "I.D.1.03"): (
        "Met", "Reviewer lists the risk-sharing mechanisms the AI missed (reinsurance, corridor, stop-loss)."),
    ("D. Payment", "I.D.2.01"): (
        "Not Met", "State cites 42 CFR 422.208/210 (Medicare Advantage) rather than 438.6, as in the "
                   "surrounding items Lynn says should be unmet."),
}

OUTCOME_FILLS = {
    "Improved": "C6EFCE",
    "Limited": "FFC7CE",
    "Still correct": "E2EFDA",
    "Still incorrect": "FFEB9C",
}


def status_label(value):
    """Canonical display form, e.g. "Not Met", or "Blank"/"Unrecognized" from normalize_status."""
    return normalize_status(value)


def infer_expected(sheet_title, item, feedback, old_status):
    override = EXPECTED_OVERRIDES.get((sheet_title, item))
    if override:
        return override

    for pattern, status in FEEDBACK_RULES:
        match = pattern.search(feedback)
        if match:
            if status == "KEEP":
                return old_status, f"Reviewer endorsed the old AI status (\"{match.group(0)}\")."
            return status, f"Feedback says \"{match.group(0)}\"."

    # No status stated or implied: the row can't be graded, so the caller leaves it out of the
    # report. Add an EXPECTED_OVERRIDES entry to bring it in.
    return None, None


def outcome(expected, old_status, new_status):
    accepted = [status.strip() for status in expected.split(" or ")]
    old_ok, new_ok = old_status in accepted, new_status in accepted
    if new_ok:
        return "Still correct" if old_ok else "Improved"
    return "Limited" if old_ok else "Still incorrect"


def reviewer_columns(sheet):
    columns = {}
    header = next(sheet.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, max_col=MAX_SCAN_COL,
                                  values_only=True), ())
    for column, value in enumerate(header, 1):
        match = REVIEWER_HEADER.match(_text(value))
        if match:
            columns[column] = match.group(1).title()
    return columns


def _rows(sheet, max_col):
    return list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or MAX_SCAN_ROW, MAX_SCAN_ROW),
                                max_col=max_col, values_only=True))


def _cell(rows, row, column):
    if row - 1 >= len(rows):
        return ""
    values = rows[row - 1]
    return _text(values[column - 1]) if column - 1 < len(values) else ""


def collect_feedback(feedback_path, reviewed_path):
    # read_only: these workbooks carry sheets whose used range runs to XFD1048576, which makes a
    # full load (and anything that walks ws.iter_rows() unbounded) take minutes.
    feedback_wb = load_workbook(feedback_path, read_only=True)
    reviewed_wb = load_workbook(reviewed_path, read_only=True)

    records, mismatched, ungraded = [], [], []
    for sheet in feedback_wb.worksheets:
        columns = reviewer_columns(sheet)
        if not columns:
            continue
        if sheet.title not in reviewed_wb.sheetnames:
            log.info(f"collect_feedback() {sheet.title} has reviewer feedback but is missing from "
                     f"{reviewed_path.name}, skipped")
            continue

        old_rows = _rows(sheet, max(columns))
        new_rows = _rows(reviewed_wb[sheet.title], COMMENTS_COL)

        for row in range(FIRST_DATA_ROW, len(old_rows) + 1):
            for column, reviewer in columns.items():
                feedback = _cell(old_rows, row, column)
                if not feedback:
                    continue

                item = _cell(old_rows, row, ITEM_COL)
                new_item = _cell(new_rows, row, ITEM_COL)
                if item != new_item:
                    mismatched.append((sheet.title, row, item, new_item))

                old_status = status_label(_cell(old_rows, row, STATUS_COL))
                new_status = status_label(_cell(new_rows, row, STATUS_COL))
                expected, basis = infer_expected(sheet.title, item, feedback, old_status)
                if expected is None:
                    ungraded.append((sheet.title, row, item, reviewer))
                    continue

                records.append({
                    "values": [
                        sheet.title, row, item,
                        _cell(old_rows, row, LEGAL_CITE_COL), _cell(old_rows, row, REQUIREMENT_COL),
                        old_status, _cell(old_rows, row, WHERE_FOUND_COL),
                        _cell(old_rows, row, FOLLOW_UP_COL), _cell(old_rows, row, COMMENTS_COL),
                        reviewer, feedback, expected,
                        new_status, _cell(new_rows, row, WHERE_FOUND_COL),
                        _cell(new_rows, row, FOLLOW_UP_COL), _cell(new_rows, row, COMMENTS_COL),
                        "Yes" if old_status != new_status else "No",
                        outcome(expected, old_status, new_status),
                    ],
                    "basis": basis,
                })

    log.info(f"collect_feedback() Collected {len(records)} reviewer feedback rows from {feedback_path.name}")
    return records, mismatched, ungraded


def build_summary(workbook, records):
    sheet = workbook.create_sheet("Summary")
    reviewers = sorted({record["values"][9] for record in records})
    outcomes = list(OUTCOME_FILLS)

    sheet.append(["Outcome vs Feedback"] + reviewers + ["Total"])
    for name in outcomes:
        counts = [sum(1 for record in records
                      if record["values"][9] == reviewer and record["values"][17] == name)
                  for reviewer in reviewers]
        sheet.append([name] + counts + [sum(counts)])
    totals = [sum(1 for record in records if record["values"][9] == reviewer) for reviewer in reviewers]
    sheet.append(["Total"] + totals + [sum(totals)])

    changed = sum(1 for record in records if record["values"][16] == "Yes")
    sheet.append([])
    sheet.append(["Status changed (old -> new)", changed])

    for row in sheet.iter_rows(min_row=1, max_row=len(outcomes) + 2):
        for cell in row:
            cell.border = THIN_BORDER
            if cell.row == 1 or cell.row == len(outcomes) + 2:
                cell.font = Font(bold=True)
    sheet.column_dimensions["A"].width = 28
    return sheet


def build_report(records, report_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Feedback Comparison"
    sheet.append(HEADER)

    expected_col = HEADER.index("Expected Status (from feedback)") + 1
    outcome_col = HEADER.index("Outcome vs Feedback") + 1

    for record in records:
        sheet.append(record["values"])
        row = sheet.max_row
        sheet.cell(row=row, column=expected_col).comment = Comment(record["basis"], "compare_feedback")
        fill = OUTCOME_FILLS.get(record["values"][-1])
        if fill:
            sheet.cell(row=row, column=outcome_col).fill = PatternFill("solid", fgColor=fill)

    for column, width in enumerate(WIDTHS, 1):
        sheet.column_dimensions[sheet.cell(row=1, column=column).column_letter].width = width
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=len(HEADER)):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.row == 1:
                cell.font = Font(bold=True)
    sheet.freeze_panes = "F2"
    sheet.auto_filter.ref = sheet.dimensions

    build_summary(workbook, records)
    workbook.save(report_path)
    log.info(f"build_report() Wrote feedback comparison report to {report_path}")
    return report_path


def run(args):
    folder = default_folder()
    feedback_path = Path(args.feedback) if Path(args.feedback).is_absolute() else folder / args.feedback
    reviewed_path = Path(args.reviewed) if Path(args.reviewed).is_absolute() else folder / args.reviewed
    for path in (feedback_path, reviewed_path):
        if not path.is_file():
            print(f"Workbook not found: {path}")
            return 1

    records, mismatched, ungraded = collect_feedback(feedback_path, reviewed_path)
    if not records:
        print(f"No reviewer feedback columns (\"<Name>'s comments\" on row {HEADER_ROW}) found in "
              f"{feedback_path.name}")
        return 1

    for title, row, item, new_item in mismatched:
        print(f"  WARNING {title} row {row}: item {item!r} in feedback workbook vs {new_item!r} in reviewed")
    for title, row, item, reviewer in ungraded:
        print(f"  Left out {title} row {row} ({item}): {reviewer}'s feedback states no expected status; "
              f"add it to EXPECTED_OVERRIDES to include it")

    report_path = Path(args.out) if args.out else reviewed_path.with_name(REPORT_NAME)
    if not saveable(report_path):
        print(f"{report_path} is open in another program (likely Excel); close it and re-run")
        return 1
    build_report(records, report_path)

    tally = {}
    for record in records:
        key = (record["values"][9], record["values"][-1])
        tally[key] = tally.get(key, 0) + 1
    print(f"{len(records)} feedback rows compared")
    for (reviewer, name), count in sorted(tally.items()):
        print(f"  {reviewer:<8}{name:<16}{count:>4}")
    print(f"\nWritten feedback comparison report to {report_path}")
    return 0


def main():
    sys.stdout.reconfigure(errors="replace")

    parser = argparse.ArgumentParser(description=DESCRIPTION)
    parser.add_argument("--feedback", required=True,
                        help="Earlier reviewed workbook with reviewer comment columns (path or name in output_excel/)")
    parser.add_argument("--reviewed", required=True,
                        help="Newer .reviewed workbook to compare against (path or name in output_excel/)")
    parser.add_argument("--out", help=f"Report path (default: {REPORT_NAME} next to --reviewed)")
    args = parser.parse_args()

    return run(args)


if __name__ == "__main__":
    sys.exit(main())
